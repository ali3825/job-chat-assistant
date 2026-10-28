"""Resume parsing: PDF/DOCX -> text -> structured Profile (rule-based, no LLM, no network).

Every extracted field is tagged in ``Profile.provenance`` as ``stated`` (matched
directly in the text) or ``inferred`` (guessed from position/dates). Missing
fields stay None/empty -- nothing is invented.
"""
from __future__ import annotations

import io
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import date
from pathlib import Path

from ..core.skills import extract_skills
from ..schemas import EducationItem, Profile, ProjectItem

MIN_TEXT_CHARS = 80
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024


class ResumeError(Exception):
    def __init__(self, code: str, message: str, recovery: str, status: int = 422):
        super().__init__(message)
        self.code, self.message, self.recovery, self.status = code, message, recovery, status


# --------------------------------------------------------------------------- text extraction

def validate_upload(filename: str, data: bytes, max_bytes: int) -> str:
    """Return the normalised extension ('pdf' | 'docx') or raise ResumeError."""
    if not data:
        raise ResumeError("empty_file", "The uploaded file is empty.", "Choose a resume file and try again.", 400)
    if len(data) > max_bytes:
        raise ResumeError("file_too_large", f"The file is larger than {max_bytes // (1024 * 1024)} MB.",
                          "Upload a smaller PDF or DOCX.", 413)
    ext = Path(filename or "").suffix.lower().lstrip(".")
    if ext not in {"pdf", "docx"}:
        raise ResumeError("unsupported_type", "Only PDF and DOCX resumes are supported.",
                          "Save your resume as a PDF or DOCX and upload it again.", 415)
    # Check magic bytes rather than trusting the extension/content-type.
    if ext == "pdf" and not data.lstrip()[:5] == b"%PDF-":
        raise ResumeError("corrupt_file", "This file has a .pdf name but is not a valid PDF.",
                          "Re-export the resume as a PDF and try again.")
    if ext == "docx" and data[:4] != b"PK\x03\x04":
        raise ResumeError("corrupt_file", "This file has a .docx name but is not a valid DOCX.",
                          "Re-save the resume as DOCX (not .doc) and try again.")
    return ext


def extract_text(data: bytes, ext: str, max_pages: int = 20) -> tuple[str, list[str]]:
    warnings: list[str] = []
    if ext == "pdf":
        text = _pdf_text(data, max_pages, warnings)
    else:
        text = _docx_text(data)
    text = _clean(text)
    if len(text) < MIN_TEXT_CHARS:
        raise ResumeError(
            "unreadable",
            "I couldn't read any usable text from this resume.",
            "If it's a scan or image, upload a text-based PDF/DOCX, or tell me your skills and "
            "experience in chat (e.g. “add skill Python”, “I have 2 years of experience”).",
        )
    return text, warnings


def _pdf_text(data: bytes, max_pages: int, warnings: list[str]) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                if reader.decrypt("") == 0:
                    raise ResumeError("encrypted", "This PDF is password-protected.",
                                      "Remove the password and upload it again.")
            except ResumeError:
                raise
            except Exception:
                raise ResumeError("encrypted", "This PDF is password-protected.",
                                  "Remove the password and upload it again.")
        pages = reader.pages[:max_pages]
        if len(reader.pages) > max_pages:
            warnings.append(f"Only the first {max_pages} pages were read.")
        text = "\n".join((p.extract_text() or "") for p in pages)
    except ResumeError:
        raise
    except (PyPdfError, ValueError, KeyError, OSError, TypeError, AttributeError, RecursionError):
        raise ResumeError("corrupt_file", "This PDF appears to be damaged and can't be opened.",
                          "Re-export the resume as a PDF and try again.")
    if len(_clean(text)) < MIN_TEXT_CHARS:
        ocr = _ocr_pdf(data)
        if ocr:
            warnings.append("This looked like a scanned PDF, so I used OCR; please double-check the details.")
            return ocr
    return text


def _ocr_pdf(data: bytes) -> str:
    """OCR fallback using the pdftoppm + tesseract binaries when installed. Returns '' if unavailable."""
    if not (shutil.which("pdftoppm") and shutil.which("tesseract")):
        return ""
    try:
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "in.pdf"
            pdf.write_bytes(data)
            subprocess.run(["pdftoppm", "-r", "150", "-l", "3", "-png", str(pdf), str(Path(tmp) / "p")],
                           check=True, timeout=60, capture_output=True)
            out = []
            for img in sorted(Path(tmp).glob("p*.png")):
                r = subprocess.run(["tesseract", str(img), "stdout"], check=True, timeout=60,
                                   capture_output=True, text=True)
                out.append(r.stdout)
            return "\n".join(out)
    except (subprocess.SubprocessError, OSError):
        return ""


def _docx_text(data: bytes) -> str:
    from docx import Document

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
            if len(infos) > 2000 or sum(i.file_size for i in infos) > MAX_DOCX_UNCOMPRESSED:
                raise ResumeError("corrupt_file", "This DOCX file is not accepted (unusually large contents).",
                                  "Re-save the resume as a normal DOCX and try again.")
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    except ResumeError:
        raise
    except Exception:
        raise ResumeError("corrupt_file", "This DOCX file appears to be damaged and can't be opened.",
                          "Re-save the resume as DOCX and try again.")


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("•", "•")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# --------------------------------------------------------------------------- structured parsing

SECTION_PATTERNS = {
    "skills": r"(?:technical\s+|key\s+)?skills(?:\s+(?:&|and)\s+tools)?|core competencies|technologies|tech stack|tools(?:\s*&\s*technologies)?",
    "experience": r"(?:work\s+|professional\s+|relevant\s+)?experience|employment(?:\s+history)?|work history|internships?",
    "education": r"education(?:al)?(?:\s+(?:background|qualifications?))?|academics?|academic\s+(?:background|qualifications?)",
    "projects": r"(?:personal\s+|academic\s+|key\s+|selected\s+)?projects",
    "certifications": r"certifications?|certificates?|licenses?(?:\s*&\s*certifications?)?|courses",
    "summary": r"(?:professional\s+)?summary|objective|career objective|profile|about me",
}
_SECTION_RES = {k: re.compile(rf"^\s*(?:{v})\s*:?\s*$", re.IGNORECASE) for k, v in SECTION_PATTERNS.items()}

TITLE_WORDS = (r"engineer|developer|analyst|designer|scientist|manager|intern|consultant|architect|"
               r"administrator|tester|qa|devops|programmer|specialist|lead|associate|officer|researcher|sde")
_TITLE_RE = re.compile(rf"\b(?:{TITLE_WORDS})\b", re.IGNORECASE)
_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
_PHONE_RE = re.compile(r"(?<!\d)(\+?\d[\d\s().-]{8,16}\d)(?!\d)")
_LINK_RE = re.compile(r"(?:https?://)?(?:www\.)?(?:linkedin\.com/in/|github\.com/)[A-Za-z0-9_./-]+", re.IGNORECASE)
_MONTHS = "jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec"
_DATE = rf"(?:(?:{_MONTHS})[a-z]*\.?\s+\d{{4}}|\d{{1,2}}/\d{{4}}|\d{{4}})"
_RANGE_RE = re.compile(rf"({_DATE})\s*(?:-|–|—|to)\s*({_DATE}|present|current|now|ongoing)", re.IGNORECASE)
_YEARS_RE = re.compile(r"(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)(?:\s+of)?\s+(?:[a-z]+\s+){0,2}experience", re.IGNORECASE)
_DEGREE_RE = re.compile(
    r"\b(B\.?\s?Tech|B\.?\s?E\.?|B\.?\s?Sc\.?|B\.?\s?C\.?\s?A\.?|B\.?\s?S\.?|Bachelor[’']?s?(?:\s+of\s+[A-Za-z &]+)?|"
    r"M\.?\s?Tech|M\.?\s?Sc\.?|M\.?\s?C\.?\s?A\.?|M\.?\s?B\.?\s?A\.?|M\.?\s?S\.?|Master[’']?s?(?:\s+of\s+[A-Za-z &]+)?|"
    r"Ph\.?\s?D\.?|Doctorate|Diploma(?:\s+in\s+[A-Za-z &]+)?)(?=[\s,.:;)(-]|$)")
_BULLET_RE = re.compile(r"^\s*[•●▪■◦\-–*·]\s*")


def split_sections(text: str) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        matched = None
        if len(stripped) <= 45:
            for key, rx in _SECTION_RES.items():
                if rx.match(stripped):
                    matched = key
                    break
        if matched:
            current = matched
            sections.setdefault(current, [])
        else:
            sections.setdefault(current, []).append(stripped)
    return sections


def parse_resume_text(text: str, today: date | None = None) -> Profile:
    today = today or date.today()
    sections = split_sections(text)
    header = sections.get("header", [])
    prov: dict[str, str] = {}
    profile = Profile()

    # contact
    if m := _EMAIL_RE.search(text):
        profile.email, prov["email"] = m.group(0), "stated"
    for m in _PHONE_RE.finditer("\n".join(header[:10]) or text):
        digits = re.sub(r"\D", "", m.group(1))
        if 10 <= len(digits) <= 13:
            profile.phone, prov["phone"] = m.group(1).strip(), "stated"
            break
    profile.links = list(dict.fromkeys(m.group(0) for m in _LINK_RE.finditer(text)))[:4]

    # name: first short alphabetic line near the top
    for line in header[:8]:
        if _is_name(line):
            profile.name, prov["name"] = line.strip(), "inferred"
            break

    # designation: a title-like header line, else the first title in the experience section
    for line in header[:10]:
        if line.strip() != profile.name and _looks_like_title(line):
            profile.designation, prov["designation"] = line.strip(" |-–"), "stated"
            break
    if not profile.designation:
        for line in sections.get("experience", [])[:12]:
            if _looks_like_title(line):
                profile.designation, prov["designation"] = _extract_title(line), "inferred"
                break

    # skills
    skills_text = "\n".join(sections.get("skills", []))
    skills = extract_skills(skills_text)
    for s in extract_skills(text):
        if s not in skills:
            skills.append(s)
    profile.skills = skills
    if skills:
        prov["skills"] = "stated"

    # experience years
    stated = [float(x) for x in _YEARS_RE.findall(text)]
    if stated:
        profile.years_experience, prov["years_experience"] = max(stated), "stated"
    else:
        computed = _years_from_ranges("\n".join(sections.get("experience", [])), today)
        if computed is not None:
            profile.years_experience, prov["years_experience"] = computed, "inferred"

    # education
    profile.education = _parse_education(sections.get("education", []), text)
    if profile.education:
        prov["education"] = "stated"

    # projects, certifications
    profile.projects = _parse_projects(sections.get("projects", []))
    if profile.projects:
        prov["projects"] = "stated"
    certs = [_BULLET_RE.sub("", l).strip() for l in sections.get("certifications", [])]
    profile.certifications = [c for c in certs if 3 < len(c) < 150][:10]
    if profile.certifications:
        prov["certifications"] = "stated"

    profile.provenance = prov
    return profile


def _is_name(line: str) -> bool:
    s = line.strip()
    if not s or len(s) > 40 or any(ch.isdigit() for ch in s) or "@" in s or "http" in s.lower():
        return False
    if any(rx.match(s) for rx in _SECTION_RES.values()) or _looks_like_title(s):
        return False
    words = s.replace(",", " ").split()
    return 2 <= len(words) <= 4 and all(re.fullmatch(r"[A-Za-z][A-Za-z.'’-]*", w) for w in words)


def _looks_like_title(line: str) -> bool:
    s = line.strip()
    return len(s) <= 70 and "@" not in s and bool(_TITLE_RE.search(s)) and not _RANGE_RE.fullmatch(s)


def _extract_title(line: str) -> str:
    part = re.split(r"\s+(?:at|@|\||,|-|–|—)\s+", line.strip(), maxsplit=1)[0]
    return part.strip(" |-–")[:70]


def _to_ym(token: str, today: date) -> tuple[int, int] | None:
    t = token.strip().lower()
    if t in {"present", "current", "now", "ongoing"}:
        return today.year, today.month
    if m := re.fullmatch(rf"({_MONTHS})[a-z]*\.?\s+(\d{{4}})", t):
        return int(m.group(2)), ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"].index(m.group(1)[:3]) + 1
    if m := re.fullmatch(r"(\d{1,2})/(\d{4})", t):
        return int(m.group(2)), min(max(int(m.group(1)), 1), 12)
    if re.fullmatch(r"\d{4}", t):
        return int(t), 1
    return None


def _years_from_ranges(text: str, today: date) -> float | None:
    spans: list[tuple[int, int]] = []
    for a, b in _RANGE_RE.findall(text):
        start, end = _to_ym(a, today), _to_ym(b, today)
        if not start or not end:
            continue
        s, e = start[0] * 12 + start[1], end[0] * 12 + end[1]
        if e >= s and e - s <= 12 * 50:
            spans.append((s, e))
    if not spans:
        return None
    spans.sort()
    total, (cs, ce) = 0, spans[0]
    for s, e in spans[1:]:
        if s <= ce:
            ce = max(ce, e)
        else:
            total += ce - cs
            cs, ce = s, e
    total += ce - cs
    return round(total / 12, 1)


def _parse_education(lines: list[str], full_text: str) -> list[EducationItem]:
    source = lines or full_text.splitlines()
    items: list[EducationItem] = []
    for i, line in enumerate(source):
        m = _DEGREE_RE.search(line)
        if not m:
            continue
        window = " ".join(source[i:i + 2])
        years = [int(y) for y in re.findall(r"\b(19[89]\d|20[0-4]\d)\b", window)]
        degree = re.sub(r"\s+", " ", line.strip())[:100]
        details = re.sub(r"\s+", " ", source[i + 1].strip())[:100] if i + 1 < len(source) and not _DEGREE_RE.search(source[i + 1]) else None
        items.append(EducationItem(degree=degree, details=details, year=max(years) if years else None))
        if len(items) >= 4:
            break
    return items


def _parse_projects(lines: list[str]) -> list[ProjectItem]:
    projects: list[ProjectItem] = []
    for line in lines:
        is_bullet = bool(_BULLET_RE.match(line))
        text = _BULLET_RE.sub("", line).strip()
        if not is_bullet and 3 <= len(text) <= 100:
            projects.append(ProjectItem(name=text.split(" - ")[0].split(" | ")[0].strip(" :|"), description=None))
        elif projects and text:
            prev = projects[-1]
            desc = ((prev.description + " ") if prev.description else "") + text
            prev.description = desc[:250]
        if len(projects) >= 8:
            break
    return projects
